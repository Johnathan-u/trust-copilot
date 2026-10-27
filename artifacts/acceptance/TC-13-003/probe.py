"""Acceptance probe for TC-13-003 (synthetic workbook only)."""

import io
import json
import zipfile

import openpyxl
from app.services import fallback_export
from app.services.xlsx_writer import create_export_from_questionnaire

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "S"
ws["A1"] = "Q?"
ws["A1"].number_format = "0.00%"
ws["B2"] = "=TODAY()"
ws["A2"] = "Q2?"
buf = io.BytesIO()
wb.save(buf)
out = create_export_from_questionnaire(
    buf.getvalue(),
    [
        {"id": 1, "source_location": json.dumps({"sheet": "S", "row": 1, "col": 1})},
        {"id": 2, "source_location": json.dumps({"sheet": "S", "row": 2, "col": 1})},
    ],
    {1: "=1+1", 2: "Yes"},
)
sheet_xml = zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml").decode()
loaded = openpyxl.load_workbook(io.BytesIO(out)).active
print(
    "B1 xml has <f>:",
    "<f>1+1</f>" in sheet_xml,
    "| B1 type:",
    loaded["B1"].data_type,
    "| B1 number_format:",
    loaded["B1"].number_format,
    "| B2 (orig formula) now:",
    repr(loaded["B2"].value),
)
try:
    fallback_export.create_fallback_pack_csv([{"id": 1, "text": "q"}], {1: "=1+1"})
    print("csv ok")
except Exception as exc:
    print("csv fallback error:", type(exc).__name__, exc)
