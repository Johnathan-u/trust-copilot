"""TC-06-002i: persisted export authority before rendering, replay or effects.

SQLite retains actual model FKs/indexes; PostgreSQL uses the migrated schema.
Only object-storage I/O is replaced. No auth, source, rendering, quota or ledger
method is mocked. Full-row snapshots and exact owned-ID cleanup detect writes.
"""

from __future__ import annotations

import io
import json
import secrets
from contextlib import contextmanager
from datetime import datetime
from types import SimpleNamespace

import openpyxl
import pytest
from docx import Document as DocxDocument
from sqlalchemy import create_engine, event, select, text, update
from sqlalchemy.orm import Session

from app import worker
from app.core.database import Base
from app.core.job_errors import classify, is_retryable_code, safe_error_message
from app.models import Answer, Document, ExportRecord, Job, Question, Questionnaire, User, Workspace
from app.models.job_effect import JobEffect
from app.models.workspace_quota import WorkspaceUsage
from app.services.storage import StorageClient
from tests import _safety

MODELS = (
    User,
    Workspace,
    Document,
    Questionnaire,
    Question,
    Answer,
    Job,
    ExportRecord,
    JobEffect,
    WorkspaceUsage,
)


def snapshot(db):
    return {
        model.__tablename__: {
            row.id: tuple(row) for row in db.execute(select(model.__table__).order_by(model.id))
        }
        for model in MODELS
    }


def changed(before, after):
    return {
        table: ids
        for table in before
        if (
            ids := {
                key
                for key in before[table].keys() | after[table].keys()
                if before[table].get(key) != after[table].get(key)
            }
        )
    }


def workbook_bytes():
    book = openpyxl.Workbook()
    book.active.title = "Sheet1"
    book.active["A1"] = "Question"
    book.active["A2"] = "Synthetic question"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture(
    params=(
        pytest.param("sqlite", marks=pytest.mark.unit),
        pytest.param("postgres", marks=pytest.mark.integration),
    )
)
def export_case(request, monkeypatch):
    if request.param == "postgres":
        request.getfixturevalue("migrated_db")
        from app.core.database import engine

        db = Session(engine, expire_on_commit=False)
        own_engine = False
    else:
        engine = create_engine("sqlite://")
        with engine.begin() as connection:
            connection.exec_driver_sql("PRAGMA foreign_keys=ON")
            assert connection.exec_driver_sql("PRAGMA foreign_keys").scalar() == 1
        Base.metadata.create_all(engine, tables=[m.__table__ for m in MODELS])
        db = Session(engine, expire_on_commit=False)
        own_engine = True
    baseline = snapshot(db)
    owned = {model: set() for model in MODELS}
    before_network = _safety.blocked_attempts()
    objects, storage_calls = {}, []
    store = SimpleNamespace(objects=objects, calls=storage_calls, fail_upload=False)

    def remember(session, _context):
        for row in session.new:
            if type(row) in owned:
                owned[type(row)].add(row.id)

    event.listen(db, "after_flush", remember)

    def add(model, **fields):
        row = model(**fields)
        db.add(row)
        db.flush()
        return row

    def download(_storage, bucket, key):
        storage_calls.append(("download", bucket, key))
        return objects[(bucket, key)]

    def upload(_storage, bucket, key, body, **_kwargs):
        storage_calls.append(("upload", bucket, key))
        if store.fail_upload:
            raise ConnectionError("Synthetic object-store unavailable")
        objects[(bucket, key)] = body.read() if hasattr(body, "read") else body
        return key

    def delete(_storage, bucket, key):
        storage_calls.append(("delete", bucket, key))
        objects.pop((bucket, key), None)

    monkeypatch.setattr(StorageClient, "download", download)
    monkeypatch.setattr(StorageClient, "upload", upload)
    monkeypatch.setattr(StorageClient, "delete", delete)
    try:
        suffix = secrets.token_hex(8)
        first = db.execute(text("SELECT COALESCE(MAX(id), 0) + 1 FROM workspaces")).scalar_one()
        own = add(Workspace, id=first, name="Export own", slug=f"worker-own-{suffix}")
        foreign = add(
            Workspace, id=first + 1, name="Export foreign", slug=f"worker-foreign-{suffix}"
        )

        def questionnaire(workspace, label):
            qnr = add(
                Questionnaire,
                workspace_id=workspace.id,
                filename=f"{label}.xlsx",
                storage_key=f"raw/{suffix}/{label}.xlsx",
            )
            question = add(
                Question,
                questionnaire_id=qnr.id,
                text=f"Synthetic {label} question",
                source_location=json.dumps({"sheet": "Sheet1", "row": 2, "col": 1}),
            )
            add(Answer, question_id=question.id, text=f"PRIVATE {label} ANSWER", status="approved")
            objects[(worker.settings.s3_bucket_raw, qnr.storage_key)] = workbook_bytes()
            return qnr

        qnr = questionnaire(own, "OWN")
        other_qnr = questionnaire(foreign, "FOREIGN")
        alternate_qnr = questionnaire(own, "OTHER-OWN")
        payload = {"workspace_id": own.id, "questionnaire_id": qnr.id, "format": "xlsx"}
        job = add(
            Job,
            workspace_id=own.id,
            kind="export",
            status="running",
            attempt=1,
            payload=json.dumps(payload),
            result=json.dumps({"prior_attempt": "kept"}),
            error="Prior error",
            error_code="connection_error",
        )
        db.commit()
        yield SimpleNamespace(
            db=db,
            engine=engine,
            add=add,
            own=own,
            foreign=foreign,
            qnr=qnr,
            other_qnr=other_qnr,
            alternate_qnr=alternate_qnr,
            job=job,
            payload=payload,
            store=store,
        )
    finally:
        db.rollback()
        for model in reversed(MODELS):
            if owned[model]:
                db.query(model).filter(model.id.in_(owned[model])).delete(synchronize_session=False)
        db.commit()
        assert snapshot(db) == baseline
        assert _safety.blocked_attempts() == before_network
        event.remove(db, "after_flush", remember)
        db.close()
        if own_engine:
            engine.dispose()


def persist_payload(case, payload):
    case.job.payload = json.dumps(payload)
    case.db.commit()


@contextmanager
def trace_sql(case):
    statements = []

    def record(_conn, _cursor, statement, _parameters, _context, _many):
        statements.append(statement)

    event.listen(case.engine, "before_cursor_execute", record)
    try:
        yield statements
    finally:
        event.remove(case.engine, "before_cursor_execute", record)


def assert_denied_without_effects(case):
    before = snapshot(case.db)
    objects = dict(case.store.objects)
    with trace_sql(case) as statements:
        with pytest.raises(ValueError) as caught:
            worker.run_job(case.job, case.db)
    # Check pending ORM writes before rollback, then durable full-row state after it.
    assert not case.db.new and not case.db.dirty and not case.db.deleted
    case.db.rollback()
    assert snapshot(case.db) == before
    assert not case.store.calls and case.store.objects == objects
    assert {s.lstrip().split(None, 1)[0].upper() for s in statements} <= {"SELECT"}
    assert classify(caught.value) == "bad_input"
    assert not is_retryable_code(classify(caught.value))
    assert safe_error_message(caught.value) == "Export references are unavailable"


@pytest.mark.parametrize(
    "case_name",
    [
        "foreign_payload_and_parent",
        "foreign_workspace",
        "foreign_parent",
        "job_moved",
        "job_moved_stale_identity",
        "job_deleted",
        "parent_moved",
        "parent_moved_stale_identity",
        "parent_archived",
        "parent_archived_stale_identity",
        "parent_deleted",
        "parent_missing",
        "parent_without_file",
        "parent_without_file_stale_identity",
    ],
)
@pytest.mark.parametrize("replay", [False, True], ids=["new-export", "replay"])
def test_current_job_payload_and_active_parent_must_agree(export_case, case_name, replay):
    c = export_case
    if replay:
        c.add(
            ExportRecord,
            workspace_id=c.own.id,
            questionnaire_id=c.qnr.id,
            job_id=c.job.id,
            storage_key="exports/OWN-private-existing.xlsx",
            filename="OWN-private.xlsx",
        )
        c.db.commit()
    if case_name == "foreign_payload_and_parent":
        c.payload.update(workspace_id=c.foreign.id, questionnaire_id=c.other_qnr.id)
    elif case_name == "foreign_workspace":
        c.payload["workspace_id"] = c.foreign.id
    elif case_name == "foreign_parent":
        c.payload["questionnaire_id"] = c.other_qnr.id
    elif case_name == "job_moved":
        c.job.workspace_id = c.foreign.id
    elif case_name == "job_moved_stale_identity":
        c.db.execute(
            update(Job).where(Job.id == c.job.id).values(workspace_id=c.foreign.id),
            execution_options={"synchronize_session": False},
        )
        assert c.job.workspace_id == c.own.id
    elif case_name == "job_deleted":
        c.db.delete(c.job)
    elif case_name.startswith("parent_moved"):
        c.db.execute(
            update(Questionnaire)
            .where(Questionnaire.id == c.qnr.id)
            .values(workspace_id=c.foreign.id),
            execution_options={
                "synchronize_session": False if case_name.endswith("stale_identity") else "fetch"
            },
        )
    elif case_name.startswith("parent_archived"):
        c.db.execute(
            update(Questionnaire)
            .where(Questionnaire.id == c.qnr.id)
            .values(deleted_at=datetime(2026, 1, 1)),
            execution_options={
                "synchronize_session": False if case_name.endswith("stale_identity") else "fetch"
            },
        )
    elif case_name == "parent_deleted":
        c.db.delete(c.qnr)
    elif case_name == "parent_missing":
        c.payload["questionnaire_id"] = 2147483647
    elif case_name == "parent_without_file":
        c.qnr.storage_key = None
    elif case_name == "parent_without_file_stale_identity":
        c.db.execute(
            update(Questionnaire).where(Questionnaire.id == c.qnr.id).values(storage_key=None),
            execution_options={"synchronize_session": False},
        )
        assert c.qnr.storage_key is not None
    persist_payload(c, c.payload)
    assert_denied_without_effects(c)


@pytest.mark.parametrize("field", ["workspace_id", "questionnaire_id"])
@pytest.mark.parametrize(
    "value",
    [
        None,
        0,
        -1,
        True,
        False,
        "own-id",
        "",
        "1e0",
        "1.0",
        "float-own-id",
        [],
        {},
        2147483648,
        10**100,
    ],
)
def test_reference_types_are_not_coerced_into_authority(export_case, field, value):
    c = export_case
    if value == "own-id":
        value = str(c.payload[field])
    elif value == "float-own-id":
        value = float(c.payload[field])
    c.payload[field] = value
    persist_payload(c, c.payload)
    assert_denied_without_effects(c)


@pytest.mark.parametrize("payload", [None, [], True, 1, "payload"])
def test_non_object_payload_is_neutral_failure(export_case, payload):
    persist_payload(export_case, payload)
    assert_denied_without_effects(export_case)


@pytest.mark.parametrize(
    "reference",
    [
        "foreign_workspace",
        "foreign_questionnaire",
        "different_own_questionnaire",
        "null_questionnaire",
        "moved_workspace_stale_identity",
        "moved_parent_stale_identity",
    ],
)
def test_replay_record_must_match_exact_current_parent_and_tenant(export_case, reference):
    c = export_case
    rec = c.add(
        ExportRecord,
        workspace_id=c.own.id,
        questionnaire_id=c.qnr.id,
        job_id=c.job.id,
        storage_key="exports/PRIVATE-REPLAY.xlsx",
        filename="PRIVATE-REPLAY.xlsx",
    )
    if reference == "foreign_workspace":
        rec.workspace_id = c.foreign.id
    elif reference == "foreign_questionnaire":
        rec.questionnaire_id = c.other_qnr.id
    elif reference == "different_own_questionnaire":
        rec.questionnaire_id = c.alternate_qnr.id
    elif reference == "null_questionnaire":
        rec.questionnaire_id = None
    c.db.commit()
    if reference.startswith("moved_"):
        values = (
            {"workspace_id": c.foreign.id}
            if "workspace" in reference
            else {"questionnaire_id": c.alternate_qnr.id}
        )
        c.db.execute(
            update(ExportRecord).where(ExportRecord.id == rec.id).values(**values),
            execution_options={"synchronize_session": False},
        )
        c.db.commit()
        assert rec.workspace_id == c.own.id and rec.questionnaire_id == c.qnr.id
    assert_denied_without_effects(c)


@pytest.mark.parametrize("format", ["xlsx", "docx"])
def test_own_export_renders_only_own_answers_and_replays_once(export_case, format):
    c = export_case
    c.payload["format"] = format
    persist_payload(c, c.payload)
    before = snapshot(c.db)
    job_before = dict(
        c.db.execute(select(Job.__table__).where(Job.id == c.job.id)).mappings().one()
    )
    worker.run_job(c.job, c.db)
    c.db.commit()
    result = json.loads(c.job.result)
    record = c.db.get(ExportRecord, result["export_record_id"])
    assert (record.workspace_id, record.questionnaire_id, record.job_id) == (
        c.own.id,
        c.qnr.id,
        c.job.id,
    )
    assert result == {"export_record_id": record.id, "storage_key": record.storage_key}
    assert record.storage_key == f"exports/{c.own.id}/job-{c.job.id}.{format}"
    body = c.store.objects[(worker.settings.s3_bucket_exports, record.storage_key)]
    if format == "xlsx":
        book = openpyxl.load_workbook(io.BytesIO(body))
        rendered = "\n".join(str(cell.value) for sheet in book for row in sheet for cell in row)
    else:
        document = DocxDocument(io.BytesIO(body))
        rendered = "\n".join(
            [p.text for p in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
    assert "PRIVATE OWN ANSWER" in rendered
    assert "PRIVATE FOREIGN ANSWER" not in rendered and "PRIVATE OTHER-OWN ANSWER" not in rendered
    effects = c.db.query(JobEffect).filter(JobEffect.job_id == c.job.id).all()
    assert {(effect.workspace_id, effect.effect, effect.state) for effect in effects} == {
        (c.own.id, "usage:exports", "applied"),
        (c.own.id, "export_record", "applied"),
    }
    usage = c.db.query(WorkspaceUsage).filter(WorkspaceUsage.workspace_id == c.own.id).one()
    assert (usage.resource_type, usage.count) == ("exports", 1)
    differences = changed(before, snapshot(c.db))
    assert differences == {
        "jobs": {c.job.id},
        "export_records": {record.id},
        "job_effects": {effect.id for effect in effects},
        "workspace_usage": {usage.id},
    }
    job_after = dict(c.db.execute(select(Job.__table__).where(Job.id == c.job.id)).mappings().one())
    assert {key for key in job_before if job_before[key] != job_after[key]} <= {
        "result",
        "updated_at",
    }
    assert [call[0] for call in c.store.calls] == ["download", "upload"]
    replay_before = snapshot(c.db)
    objects_before, calls_before = dict(c.store.objects), list(c.store.calls)
    worker.run_job(c.job, c.db)
    c.db.commit()
    assert json.loads(c.job.result) == {**result, "outcome": "already_applied"}
    assert changed(replay_before, snapshot(c.db)) == {"jobs": {c.job.id}}
    assert c.store.objects == objects_before and c.store.calls == calls_before


def test_own_storage_failure_leaves_no_artifact_or_quota_commit(export_case):
    c = export_case
    before = snapshot(c.db)
    objects = dict(c.store.objects)
    c.store.fail_upload = True
    with pytest.raises(ConnectionError):
        worker.run_job(c.job, c.db)
    c.db.rollback()
    assert snapshot(c.db) == before
    assert c.store.objects == objects
    assert [call[0] for call in c.store.calls] == ["download", "upload"]
