"""TC-09-001a: real XLSX question/heading collisions and source locations."""

import json
from pathlib import Path

from openpyxl import Workbook

from app.services.xlsx_questionnaire_parser import parse_xlsx_questionnaire


def test_xlsx_detects_seven_questions_without_replacing_sections(tmp_path: Path) -> None:
    """Numbering and uppercase text must not turn questions into section labels."""
    workbook = Workbook()
    access = workbook.active
    access.title = "Access"
    for row in [
        ["ACCESS CONTROL"],
        [None, "  DO YOU REQUIRE MFA?  "],
        [12, None, "1. Do you encrypt customer data?"],
        ["2. ACCESS MANAGEMENT"],
        [None, "Q3) DESCRIBE YOUR ACCESS REVIEW PROCESS", "Do you review accounts?"],
        [None, None, "ЧИ ЗАХИЩЕНІ ВАШІ ДАНІ?"],
        [" Do you \n review access annually? "],
    ]:
        access.append(row)
    recovery = workbook.create_sheet("Recovery")
    for row in [
        ["Describe your recovery process"],
        ["RECOVERY"],
        ["1. BACKUP AND RECOVERY"],
        [None, "1.1. DO YOU TEST RESTORES?"],
    ]:
        recovery.append(row)
    path = tmp_path / "question-heading-corpus.xlsx"
    workbook.save(path)
    workbook.close()

    questions = parse_xlsx_questionnaire(path)

    assert [question["text"] for question in questions] == [
        "do you require mfa",
        "1. do you encrypt customer data",
        "q3) describe your access review process",
        "чи захищені ваші дані",
        "do you review access annually",
        "describe your recovery process",
        "1.1. do you test restores",
    ]
    assert [question["section"] for question in questions] == [
        "ACCESS CONTROL",
        "ACCESS CONTROL",
        "2. ACCESS MANAGEMENT",
        "2. ACCESS MANAGEMENT",
        "2. ACCESS MANAGEMENT",
        None,
        "1. BACKUP AND RECOVERY",
    ]
    assert [question["answer_type"] for question in questions] == [
        "yes_no",
        "yes_no",
        "free_text",
        "free_text",
        "yes_no",
        "free_text",
        "yes_no",
    ]
    assert [json.loads(question["source_location"]) for question in questions] == [
        {"file": str(path), "sheet": "Access", "row": 2, "col": 2},
        {"file": str(path), "sheet": "Access", "row": 3, "col": 3},
        {"file": str(path), "sheet": "Access", "row": 5, "col": 2},
        {"file": str(path), "sheet": "Access", "row": 6, "col": 3},
        {"file": str(path), "sheet": "Access", "row": 7, "col": 1},
        {"file": str(path), "sheet": "Recovery", "row": 1, "col": 1},
        {"file": str(path), "sheet": "Recovery", "row": 4, "col": 2},
    ]
    assert [question["confidence"] for question in questions] == [85] * 7


def test_xlsx_numbered_and_uppercase_headings_are_not_questions(tmp_path: Path) -> None:
    """A numbering marker alone is not sufficient to override a real heading."""
    workbook = Workbook()
    sheet = workbook.active
    for heading in [
        "ACCESS CONTROL",
        "2. ACCESS MANAGEMENT",
        "1. BACKUP AND RECOVERY",
        "1.0 CONTROL ENVIRONMENT",
        "1.2. Risk Assessment",
        "Q3) ACCESS MANAGEMENT",
        "УПРАВЛІННЯ ДОСТУПОМ",
        "DOCUMENT RETENTION",
        "LIST OF POLICIES",
        "STATE LAWS",
        "ADDRESS INFORMATION",
        "1. Document Retention",
        "1.1. List of Policies",
    ]:
        sheet.append([heading])
    path = tmp_path / "headings.xlsx"
    workbook.save(path)
    workbook.close()

    assert parse_xlsx_questionnaire(path) == []
