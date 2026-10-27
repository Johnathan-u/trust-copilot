"""TC-13-003: formula-like answers are written as literal spreadsheet text.

Synthetic data only; no DB, network, storage or model calls.
Ticket proposes tests/release/test_export_literal_text.py; per Coordinator Q3
this lives outside tests/release/ (owned by TC-02-001).
"""

import csv
import io
import json
import re
import zipfile

import openpyxl
import pytest
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.formula import ArrayFormula

from app.services.fallback_export import create_fallback_pack_csv, neutralize_csv_cell
from app.services.xlsx_writer import (
    create_export_from_questionnaire,
    create_workbook,
    is_formula_like,
    place_answers_into_workbook,
    write_answers_to_sheet,
)

FORMULA_LIKE_ANSWERS = [
    "=1+1",
    "=HYPERLINK(\"http://example.invalid/x\",\"click\")",
    "+SUM(A1:A3)",
    "-2+3",
    "@SUM(A1)",
    "=cmd|' /C calc'!A0",
    "\t=1+1",
    "\r=1+1",
    " =1+1",
    "\uff1d1+1",
]
PLAIN_ANSWERS = ["Yes", "No - see policy", "123", "", "AES-256 at rest"]


def _template(extra=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Security"
    ws["A1"] = "Question"
    ws["B1"] = "Answer"
    for i in range(len(FORMULA_LIKE_ANSWERS) + len(PLAIN_ANSWERS)):
        ws.cell(row=i + 2, column=1, value=f"Synthetic question {i + 1}?")
    # Original, legitimate template formulas (must survive export).
    ws["D1"] = "=COUNTA(B2:B20)"
    ws["D2"] = "=IF(B2=\"Yes\",1,0)"
    if extra:
        extra(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _questions(n: int) -> list[dict]:
    return [
        {"id": i + 1, "source_location": json.dumps({"sheet": "Security", "row": i + 2, "col": 1})}
        for i in range(n)
    ]


def _sheet_xml(xlsx: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        return z.read("xl/worksheets/sheet1.xml").decode("utf-8")


def _cell_xml(sheet_xml: str, coord: str) -> str:
    m = re.search(rf'<c r="{coord}"[^>]*?(?:/>|>.*?</c>)', sheet_xml, re.S)
    assert m, f"cell {coord} not in sheet xml"
    return m.group(0)


def _export(answers: list[str], template: bytes | None = None, skipped=None) -> bytes:
    qs = _questions(len(answers))
    q2a = {i + 1: a for i, a in enumerate(answers)}
    return create_export_from_questionnaire(template or _template(), qs, q2a, skipped=skipped)


# ---------------- AC1: formula-like answers cannot become formulas ----------------


@pytest.mark.parametrize("answer", FORMULA_LIKE_ANSWERS)
def test_ac1_formula_like_answer_written_as_literal_string(answer: str) -> None:
    out = _export([answer])
    xml = _cell_xml(_sheet_xml(out), "B2")
    assert "<f>" not in xml and "<f " not in xml, xml
    assert 't="inlineStr"' in xml or 't="s"' in xml, xml
    ws = openpyxl.load_workbook(io.BytesIO(out), data_only=False)["Security"]
    cell = ws["B2"]
    assert cell.data_type == "s"
    assert cell.value == answer  # visible text unchanged
    assert cell.quotePrefix is True  # Excel literal-text marker


def test_ac1_no_answer_cell_has_formula_markup_mixed_batch() -> None:
    answers = FORMULA_LIKE_ANSWERS + PLAIN_ANSWERS
    out = _export(answers)
    xml = _sheet_xml(out)
    ws = openpyxl.load_workbook(io.BytesIO(out), data_only=False)["Security"]
    for i, a in enumerate(answers):
        coord = f"B{i + 2}"
        cx = _cell_xml(xml, coord) if a else ""
        assert "<f>" not in cx
        if a:
            assert ws[coord].data_type == "s"
            assert ws[coord].value == a


def test_ac1_plain_answers_unchanged_without_quote_prefix() -> None:
    out = _export(PLAIN_ANSWERS)
    ws = openpyxl.load_workbook(io.BytesIO(out))["Security"]
    for i, a in enumerate(PLAIN_ANSWERS):
        if not a:
            continue
        c = ws.cell(row=i + 2, column=2)
        assert c.value == a and c.data_type == "s"
        assert not c.quotePrefix
    # "123" stays text, not coerced to a number
    assert ws["B4"].value == "123"


def test_ac1_write_answers_to_sheet_literal() -> None:
    wb = create_workbook()
    ws = wb.active
    write_answers_to_sheet(ws, [{"row": 1, "col": 1, "text": "=2*3"}, {"row": 2, "col": 1, "text": "@x"}])
    buf = io.BytesIO()
    wb.save(buf)
    xml = _sheet_xml(buf.getvalue())
    assert "<f>" not in xml
    ws2 = openpyxl.load_workbook(io.BytesIO(buf.getvalue())).active
    assert ws2["A1"].value == "=2*3" and ws2["A1"].data_type == "s"


@pytest.mark.parametrize("value,expected", [
    ("=x", True), ("+x", True), ("-x", True), ("@x", True), ("\tx", True), ("\rx", True),
    ("  =x", True), ("\uff20x", True), ("x=1", False), ("Yes", False), ("", False), (None, False), (5, False),
])
def test_is_formula_like(value, expected) -> None:
    assert is_formula_like(value) is expected


def test_ac1_csv_fallback_neutralizes_and_returns_bytes() -> None:
    qs = [{"id": 1, "text": "=Q?"}, {"id": 2, "text": "Plain question?"}, {"id": 3, "text": "Q3"}]
    q2a = {1: "=1+1", 2: "Yes", 3: "@SUM(A1)"}
    out = create_fallback_pack_csv(qs, q2a)
    assert isinstance(out, bytes)
    rows = list(csv.reader(io.StringIO(out.decode("utf-8"))))
    assert rows[0] == ["Question", "Answer"]
    assert rows[1] == ["'=Q?", "'=1+1"]
    assert rows[2] == ["Plain question?", "Yes"]
    assert rows[3] == ["Q3", "'@SUM(A1)"]
    for r in rows[1:]:
        for v in r:
            assert not is_formula_like(v)


def test_neutralize_csv_cell_plain_passthrough() -> None:
    assert neutralize_csv_cell("Yes") == "Yes"
    assert neutralize_csv_cell(None) == ""
    assert neutralize_csv_cell("-1") == "'-1"


# ---------------- AC2: original formulas intact; formats not blindly copied ----------------


def test_ac2_original_template_formulas_preserved() -> None:
    out = _export(["=evil()", "Yes"])
    ws = openpyxl.load_workbook(io.BytesIO(out), data_only=False)["Security"]
    assert ws["D1"].data_type == "f" and ws["D1"].value == "=COUNTA(B2:B20)"
    assert ws["D2"].data_type == "f" and ws["D2"].value == '=IF(B2="Yes",1,0)'
    xml = _sheet_xml(out)
    assert "<f>COUNTA(B2:B20)</f>" in _cell_xml(xml, "D1")


def test_ac2_target_cell_with_original_formula_not_overwritten() -> None:
    def extra(ws):
        ws["B3"] = "=TODAY()"

    skipped: list[dict] = []
    out = _export(["Yes", "Overwrite attempt"], template=_template(extra), skipped=skipped)
    ws = openpyxl.load_workbook(io.BytesIO(out), data_only=False)["Security"]
    assert ws["B2"].value == "Yes"
    assert ws["B3"].data_type == "f" and ws["B3"].value == "=TODAY()"
    assert skipped == [{"sheet": "Security", "cell": "B3", "reason": "target_has_formula"}]


def test_ac2_array_formula_target_not_overwritten() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "Q?"
    ws["B1"] = ArrayFormula("B1", "=SUM(C1:C2)")
    skipped: list[dict] = []
    place_answers_into_workbook(wb, [{"sheet": "S", "row": 1, "col": 1, "text": "x"}], skipped=skipped)
    assert isinstance(ws["B1"].value, ArrayFormula)
    assert skipped and skipped[0]["cell"] == "B1"


def test_ac2_number_format_and_protection_not_copied_from_question() -> None:
    def extra(ws):
        q = ws["A2"]
        q.font = Font(bold=True)
        q.number_format = "0.00%"
        q.protection = Protection(locked=True, hidden=True)
        q.alignment = Alignment(wrap_text=True)

    out = _export(["0.5"], template=_template(extra))
    ws = openpyxl.load_workbook(io.BytesIO(out))["Security"]
    a = ws["B2"]
    assert a.value == "0.5" and a.data_type == "s"
    assert a.font.bold is True  # EXP-04 visual formatting still carried over
    assert a.alignment.wrap_text is True
    assert a.number_format == "General"
    assert a.protection.hidden is False


def test_ac2_existing_target_formatting_preserved_not_overwritten() -> None:
    def extra(ws):
        q = ws["A2"]
        q.font = Font(bold=True, color="FFFF0000")
        q.fill = PatternFill("solid", fgColor="FF00FF00")
        t = ws["B2"]
        t.font = Font(italic=True)
        t.number_format = "@"

    out = _export(["=1+1"], template=_template(extra))
    ws = openpyxl.load_workbook(io.BytesIO(out))["Security"]
    a = ws["B2"]
    assert a.font.italic is True and not a.font.bold
    assert a.fill.fill_type in (None, "none")
    assert a.number_format == "@"
    assert a.value == "=1+1" and a.data_type == "s" and a.quotePrefix is True
